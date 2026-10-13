"""Strict two-sheet workbook import; demonstration files are never modified."""
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from .engine import STAGES, dt, validate, snapshot

CASE_COLUMNS = ['Case ID', 'Request date', 'Expected delivery date', 'Partner/location', 'Current status']
EVENT_COLUMNS = ['Case ID', 'Stage', 'Event timestamp', 'Event type', 'Operational note']

def timestamp(value):
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, str):
        result = dt(value.strip())
    else:
        raise ValueError('Dates must be Excel timestamps or ISO date/time text')
    if result.tzinfo is not None:
        raise ValueError('Use consistent local timestamps without timezone offsets')
    return result.isoformat()

def records(sheet, required):
    headers = [c.value for c in sheet[1]]
    if len(headers) != len(set(headers)) or any(h not in headers for h in required):
        raise ValueError(f'{sheet.title}: missing or duplicated required columns')
    rows = []
    for cells in sheet.iter_rows(min_row=2):
        if not any(c.value is not None for c in cells):
            continue
        if any(c.data_type == 'f' for c in cells):
            raise ValueError(f'{sheet.title}: formulas are not accepted as source records')
        row = dict(zip(headers, [c.value for c in cells]))
        if any(row[h] is None for h in required if h != 'Operational note'):
            raise ValueError(f'{sheet.title}: incomplete required fields at row {cells[0].row}')
        rows.append(row)
        if len(rows) > 10000:
            raise ValueError('Maximum 10,000 rows per sheet')
    return rows

def import_workbook(blob):
    if len(blob) > 5*1024*1024:
        raise ValueError('Workbook must be smaller than 5 MB')
    try:
        with ZipFile(BytesIO(blob)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 25*1024*1024:
                raise ValueError('Expanded workbook exceeds 25 MB')
        workbook = load_workbook(BytesIO(blob), read_only=True, data_only=False)
        if set(workbook.sheetnames) != {'Cases', 'Stage Events'}:
            raise ValueError('Workbook must contain exactly Cases and Stage Events')
        raw_cases = records(workbook['Cases'], CASE_COLUMNS)
        raw_events = records(workbook['Stage Events'], EVENT_COLUMNS)
        workbook.close()
        if not raw_cases or not raw_events:
            raise ValueError('Both sheets need records')
        cases, events = [], []
        status_by_id = {}
        for row in raw_cases:
            case_id = str(row['Case ID']).strip()
            status = str(row['Current status']).strip().lower()
            if not case_id or status not in ('pending', 'completed', 'delayed', 'interrupted'):
                raise ValueError('Invalid case ID or current status')
            status_by_id[case_id] = status
            cases.append(dict(case_id=case_id, request_date=timestamp(row['Request date']),
                              expected_delivery_date=timestamp(row['Expected delivery date']),
                              partner_location=str(row['Partner/location']).strip(), pathway='Uploaded diagnostic pathway'))
        # Stable sort uses workbook row order for ties; entered/completed order must be explicit.
        for sequence, row in enumerate(raw_events):
            kind, stage = str(row['Event type']).strip().lower(), str(row['Stage']).strip()
            if kind not in ('entered', 'completed') or stage not in STAGES:
                raise ValueError('Unknown stage or event type')
            event = dict(case_id=str(row['Case ID']).strip(), stage=stage, event_type=kind,
                         timestamp=timestamp(row['Event timestamp']), sequence=sequence*2)
            events.append(event)
            if row.get('Operational note'):
                events.append(dict(event, event_type='note', sequence=sequence*2+1,
                                   category='uploaded_exception', note=str(row['Operational note'])))
        # Notes attached to a completion describe that interval; order them before completion.
        for event in events:
            if event['event_type'] == 'note':
                original = next(e for e in events if e['sequence'] == event['sequence']-1)
                if original['event_type'] == 'completed':
                    event['sequence'], original['sequence'] = original['sequence'], event['sequence']
        events.sort(key=lambda e: (e['timestamp'], e['sequence']))
        # Engine sequence expresses chronological history independently of worksheet grouping.
        for i, event in enumerate(events):
            event['sequence'] = i
        cutoff = max(e['timestamp'] for e in events)
        validate(cases, events, STAGES, cutoff)
        rows = snapshot(cases, events, cutoff)
        for row in rows:
            status = status_by_id[row['case_id']]
            if (status == 'completed') != bool(row['actual_delivery_date']):
                raise ValueError(f"{row['case_id']}: current status contradicts delivery events")
            if status == 'delayed' and not row['overdue']:
                raise ValueError(f"{row['case_id']}: delayed status conflicts with deadline at dataset cutoff")
            if status == 'interrupted':
                events.append(dict(case_id=row['case_id'], stage=row['current_stage'], timestamp=cutoff,
                                   sequence=len(events), event_type='note', status='interrupted', category='status_declaration',
                                   note='Workbook declares interrupted status at import reference cutoff; no earlier interruption time supplied.'))
        validate(cases, events, STAGES, cutoff)
        return cases, events, cutoff
    except (AssertionError, KeyError, TypeError, BadZipFile) as exc:
        raise ValueError(str(exc) or 'Invalid or contradictory stage history') from exc
